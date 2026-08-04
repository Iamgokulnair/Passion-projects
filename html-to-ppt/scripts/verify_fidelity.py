#!/usr/bin/env python3
"""Stage 4 -- content-fidelity round-trip check.

Ground truth is the slide plan (whichever element_ids the agent's Step 1 reasoning, or
the default heuristic planner, decided belong in the deck) -- NOT literally every DOM
text node in the raw manifest. Deliberately excluding page chrome (nav bars, cookie
banners, repeated boilerplate) is a legitimate Step-1 content-mapping decision, not data
loss; this check verifies that whatever the plan DID include survived intact.

Independently re-opens the saved .pptx (a fresh read, not reusing the builder's in-memory
objects) and diffs extracted text against what build_pptx.py's content_index says it wrote,
classifying every entry into one of three buckets:
  - dropped        : expected text not found anywhere on its slide -> hard fail
  - truncated       : found text is a shorter prefix of what was expected -> hard fail
  - reflowed (ok)  : text matches once whitespace/line-breaks are normalized -> informational

Usage:
    verify_fidelity.py <deck.pptx> <build_report.json> [--out fidelity_report.json]
"""
import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation


def normalize(text):
    if text is None:
        return ""
    # "|" is this pipeline's own table-cell join separator (see build_pptx.py's
    # content_index), not part of any real source text -- strip it like whitespace so
    # both sides of the diff compare on content, not on which separator each side used.
    text = text.replace("|", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def extract_slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        try:
            # Pictures only: python-pptx (through at least 1.0.x) exposes no public
            # alt-text API, so read the same descr attribute build_pptx.py writes to.
            descr = shape._element.nvPicPr.cNvPr.get("descr")  # noqa: SLF001
            if descr:
                parts.append(descr)
        except AttributeError:
            pass
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return normalize(" ".join(parts))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("build_report", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    build_report = json.loads(args.build_report.read_text(encoding="utf-8"))
    content_index = build_report.get("content_index", {})

    prs = Presentation(str(args.pptx))
    slide_texts = [extract_slide_text(s) for s in prs.slides]

    dropped, truncated, reflowed_ok, no_text_expected = [], [], [], []

    for element_id, entry in content_index.items():
        expected = entry.get("text_written")
        slide_index = entry.get("slide_index")
        if expected is None:
            no_text_expected.append(element_id)  # e.g. an image with no alt text
            continue
        expected_norm = normalize(expected)
        if not expected_norm:
            continue
        if slide_index is None or slide_index >= len(slide_texts):
            dropped.append({"element_id": element_id, "expected": expected, "reason": "no slide index recorded"})
            continue
        haystack = slide_texts[slide_index]
        if expected_norm in haystack:
            reflowed_ok.append(element_id)
            continue
        # Truncation check: does a meaningful prefix of the expected text appear?
        found_prefix = False
        for cut in (0.9, 0.75, 0.5):
            prefix = expected_norm[: int(len(expected_norm) * cut)]
            if prefix and prefix in haystack:
                truncated.append({"element_id": element_id, "expected": expected, "prefix_match_ratio": cut})
                found_prefix = True
                break
        if not found_prefix:
            dropped.append({"element_id": element_id, "expected": expected, "slide_index": slide_index})

    verdict = "PASS" if not dropped and not truncated else "FAIL"
    report = {
        "pptx": str(args.pptx),
        "verdict": verdict,
        "counts": {
            "verified_ok": len(reflowed_ok),
            "dropped": len(dropped),
            "truncated": len(truncated),
            "non_text_elements": len(no_text_expected),
        },
        "dropped": dropped,
        "truncated": truncated,
    }
    out_path = args.out or args.pptx.with_suffix(".fidelity_report.json")
    out_path.write_text(json.dumps(report, indent=2), encoding="utf-8")

    print("fidelity: " + verdict + "  (" + str(len(reflowed_ok)) + " ok, " +
          str(len(dropped)) + " dropped, " + str(len(truncated)) + " truncated)")
    print("fidelity report: " + str(out_path))
    if verdict == "FAIL":
        sys.exit(1)


if __name__ == "__main__":
    main()
