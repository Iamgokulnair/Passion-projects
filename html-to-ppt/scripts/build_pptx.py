#!/usr/bin/env python3
"""Stage 2 -- content manifest (+ optional slide plan) -> a real .pptx.

If no slide plan is given, a default heuristic planner groups the manifest's content
elements into slides at h1/h2 heading boundaries, sorted by on-page vertical position
(box.y) since parse_html.py's element ids are assigned pass-by-type, not document order.
This default plan is always written to disk (never only held in memory) so it is
inspectable and reusable as the ground truth for the fidelity gate either way -- whether
authored by an agent's Step 1 reasoning or by this fallback.

Usage:
    build_pptx.py <manifest.json> [--plan slide_plan.json] [--out deck.pptx]
"""
import argparse
import base64
import json
import re
import sys
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, str(Path(__file__).parent))
from lib import SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, safe_font  # noqa: E402

BACKGROUND_SHAPE_NAME = "html-to-ppt:background"
MARGIN_IN = 0.6
TITLE_TOP_IN = 0.4
TITLE_HEIGHT_IN = 0.9
BODY_TOP_IN = 1.5
BODY_HEIGHT_IN = SLIDE_HEIGHT_IN - BODY_TOP_IN - 0.5
CONTENT_TYPES = {"heading", "text", "list_item"}
HEADING_BREAK_LEVELS = {1, 2}


# --------------------------------------------------------------------------- planning
def default_plan(manifest):
    elements = manifest["elements"]
    by_id = {e["id"]: e for e in elements}

    content = [e for e in elements if e["type"] in CONTENT_TYPES]

    def y_of(e):
        box = (e.get("style") or {}).get("box") or {}
        return box.get("y", 0), box.get("x", 0)

    content.sort(key=y_of)

    # A new group starts when a heading-level break arrives AND the current group has
    # already claimed a title -- NOT when it merely has non-heading body content yet.
    # (A heading immediately followed by another heading, with no body text between them --
    # e.g. a lone table right under its caption -- must still split into two groups.)
    groups = []
    current = None
    for e in content:
        is_break = e["type"] == "heading" and (e.get("level") in HEADING_BREAK_LEVELS)
        starts_new = is_break and current is not None and current["title_element_id"] is not None
        if current is None or starts_new:
            current = {"title_element_id": None, "title": None, "element_ids": []}
            groups.append(current)
        if is_break and current["title_element_id"] is None:
            current["title_element_id"] = e["id"]
            current["title"] = e["text"]
        else:
            current["element_ids"].append(e["id"])
    if not groups:
        groups = [{"title_element_id": None, "title": None, "element_ids": [e["id"] for e in content]}]

    images = [e for e in elements if e["type"] == "image"]
    tables = [e for e in elements if e["type"] == "table"]
    backgrounds = [e for e in elements if e["type"] == "background_panel"]

    def start_y(g):
        eid = g["title_element_id"] or (g["element_ids"][0] if g["element_ids"] else None)
        if eid is None:
            return float("inf")
        box = (by_id[eid].get("style") or {}).get("box") or {}
        return box.get("y", float("inf"))

    # Range-based, not content-derived: each slide owns everything between its own start
    # and the NEXT slide's start, so a heading with no body text before the next heading
    # (e.g. a lone table under its caption) still correctly claims that table/image/bg.
    starts = [start_y(g) for g in groups]
    ranges = [(starts[i], starts[i + 1] if i + 1 < len(groups) else float("inf")) for i in range(len(groups))]

    slides = []
    for idx, g in enumerate(groups):
        lo, hi = ranges[idx]

        for img in images:
            iy = (img.get("style") or {}).get("box", {}).get("y", -1)
            if lo <= iy < hi and img["id"] not in g["element_ids"]:
                g["element_ids"].append(img["id"])
        for tbl in tables:
            ty = (tbl.get("style") or {}).get("box", {}).get("y", -1)
            if lo <= ty < hi and tbl["id"] not in g["element_ids"]:
                g["element_ids"].append(tbl["id"])

        # Dominant background panel: largest-area panel whose own range contains this slide's start.
        best_bg, best_area = None, 0
        for bg in backgrounds:
            box = (bg.get("style") or {}).get("box") or {}
            by0, bh = box.get("y", -1), box.get("h", 0)
            if by0 <= lo <= by0 + bh:
                area = box.get("w", 0) * bh
                if area > best_area:
                    best_bg, best_area = bg["id"], area

        has_table = any(by_id[eid]["type"] == "table" for eid in g["element_ids"])
        has_image = any(by_id[eid]["type"] == "image" for eid in g["element_ids"])
        n_list = sum(1 for eid in g["element_ids"] if by_id[eid]["type"] == "list_item")
        if has_table:
            layout = "table"
        elif has_image and len(g["element_ids"]) <= 3:
            layout = "image_feature"
        elif n_list >= 2:
            layout = "bullet_stack"
        elif not g["element_ids"]:
            layout = "section_divider"
        else:
            layout = "text_block"

        slides.append({
            "index": idx,
            "layout": layout,
            "title": g["title"],
            "title_element_id": g["title_element_id"],
            "element_ids": g["element_ids"],
            "background_element_id": best_bg,
        })

    return {"slides": slides}


# --------------------------------------------------------------------------- helpers
def hex_to_rgbcolor(hexstr):
    return RGBColor.from_string(hexstr)


def apply_background(slide, prs, bg_element, log):
    if bg_element is None:
        return
    style = bg_element.get("style") or {}
    bg = style.get("background") or {"kind": "none"}
    kind = bg.get("kind")
    if kind == "none":
        return
    left, top = Emu(0), Emu(0)
    width, height = prs.slide_width, prs.slide_height
    if kind == "solid":
        rect = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE = 1
        rect.name = BACKGROUND_SHAPE_NAME
        rect.fill.solid()
        rect.fill.fore_color.rgb = hex_to_rgbcolor(bg["hex"])
        rect.line.fill.background()
        rect.shadow.inherit = False
        _send_to_back(slide, rect)
    elif kind == "linear-gradient":
        rect = slide.shapes.add_shape(1, left, top, width, height)
        rect.name = BACKGROUND_SHAPE_NAME
        fill = rect.fill
        fill.gradient()
        try:
            fill.gradient_angle = bg.get("angle_deg", 180)
        except Exception:  # noqa: BLE001 -- angle not settable on this python-pptx version
            pass
        stops = bg.get("stops", [])
        gs = list(fill.gradient_stops)
        if len(stops) >= 2 and gs:
            gs[0].color.rgb = hex_to_rgbcolor(stops[0]["hex"])
            gs[0].position = stops[0]["pos"]
            gs[-1].color.rgb = hex_to_rgbcolor(stops[-1]["hex"])
            gs[-1].position = stops[-1]["pos"]
            if len(stops) > 2:
                log.append({
                    "element_id": bg_element["id"],
                    "note": str(len(stops)) + "-stop gradient simplified to first/last stop pair "
                            "(python-pptx's stable gradient API only exposes the stop count baked "
                            "into the preset fill, not arbitrary N stops) -- middle stops interpolate "
                            "rather than matching the source exactly. Rasterize for exact fidelity.",
                })
        rect.line.fill.background()
        rect.shadow.inherit = False
        _send_to_back(slide, rect)
    elif kind == "complex":
        asset = bg.get("asset")
        if asset and Path(asset).exists():
            pic = slide.shapes.add_picture(asset, left, top, width=width, height=height)
            pic.name = BACKGROUND_SHAPE_NAME
            _send_to_back(slide, pic)
        else:
            log.append({"element_id": bg_element["id"], "note": "background rasterization asset missing -- background left blank."})


def _send_to_back(slide, shape):
    spTree = slide.shapes._spTree
    spTree.remove(shape._element)
    spTree.insert(2, shape._element)  # after nvGrpSpPr/grpSpPr, before every other shape


def resolve_image_bytes(element, images_cache_dir, log):
    src = element.get("src_resolved") or element.get("src") or ""
    if src.startswith("data:image"):
        try:
            header, b64data = src.split(",", 1)
            ext = "png" if "png" in header else ("jpg" if "jpeg" in header or "jpg" in header else "png")
            images_cache_dir.mkdir(parents=True, exist_ok=True)
            path = images_cache_dir / (element["id"] + "." + ext)
            path.write_bytes(base64.b64decode(b64data))
            return path
        except Exception as err:  # noqa: BLE001
            log.append({"element_id": element["id"], "note": "data: URI decode failed: " + str(err)})
            return None
    if src.startswith("file://"):
        path = Path(src[len("file://"):])
        return path if path.exists() else None
    if src.startswith("http://") or src.startswith("https://"):
        try:
            images_cache_dir.mkdir(parents=True, exist_ok=True)
            ext = src.rsplit(".", 1)[-1][:4] if "." in src.rsplit("/", 1)[-1] else "png"
            path = images_cache_dir / (element["id"] + "." + ext)
            urllib.request.urlretrieve(src, path)  # noqa: S310 -- explicit remote-image fetch, documented network dependency
            return path
        except Exception as err:  # noqa: BLE001
            log.append({"element_id": element["id"], "note": "remote image fetch failed: " + str(err)})
            return None
    return None


def add_title(slide, text, font_family, log):
    box = slide.shapes.add_textbox(Inches(MARGIN_IN), Inches(TITLE_TOP_IN),
                                    Inches(SLIDE_WIDTH_IN - 2 * MARGIN_IN), Inches(TITLE_HEIGHT_IN))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    font_name, substituted = safe_font(font_family or "")
    run.font.name = font_name
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("1A1A1A")
    if substituted and font_family:
        log.append({"role": "title", "requested": font_family, "used": font_name})
    return box


def add_body_textblock(slide, lines, content_index, source_ids, font_family, log, bullets=False, region=None):
    left_in, top_in, width_in, height_in = region or (MARGIN_IN, BODY_TOP_IN, SLIDE_WIDTH_IN - 2 * MARGIN_IN, BODY_HEIGHT_IN)
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tf = box.text_frame
    tf.word_wrap = True
    font_name, substituted = safe_font(font_family or "")
    if substituted and font_family:
        log.append({"role": "body", "requested": font_family, "used": font_name})
    for i, (eid, text) in enumerate(zip(source_ids, lines)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = ("•  " + text) if bullets else text
        run.font.name = font_name
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor.from_string("2B2B2B")
        content_index[eid] = {"slide_index": None, "text_written": text}
    return box


def add_table(slide, rows, source_id, content_index, log):
    if not rows:
        return
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(MARGIN_IN), Inches(BODY_TOP_IN),
        Inches(SLIDE_WIDTH_IN - 2 * MARGIN_IN), Inches(min(BODY_HEIGHT_IN, 0.4 * n_rows + 0.2)),
    )
    table = table_shape.table
    written_cells = []
    for r, row in enumerate(rows):
        for c in range(n_cols):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
            written_cells.append(text)
    content_index[source_id] = {"slide_index": None, "text_written": " | ".join(written_cells)}


def add_image(slide, element, images_cache_dir, content_index, log, region=None):
    path = resolve_image_bytes(element, images_cache_dir, log)
    if path is None:
        content_index[element["id"]] = {"slide_index": None, "text_written": None, "placed": False}
        return
    left_in, top_in, width_in, height_in = region or (MARGIN_IN, BODY_TOP_IN, SLIDE_WIDTH_IN - 2 * MARGIN_IN, BODY_HEIGHT_IN)
    # Height-only sizing preserves the source image's aspect ratio (python-pptx scales
    # width to match); width_in is the space allotted, not a forced/distorting fill.
    pic = slide.shapes.add_picture(str(path), Inches(left_in), Inches(top_in), height=Inches(height_in))
    if pic.width > Inches(width_in):
        pic.width, pic.height = Inches(width_in), int(Inches(width_in) * pic.height / pic.width)
    alt = element.get("alt") or None
    if alt:
        # python-pptx (through at least 1.0.x) has no public alt-text API; the descr
        # attribute on cNvPr is what PowerPoint itself reads as alt text.
        pic._element.nvPicPr.cNvPr.set("descr", alt)  # noqa: SLF001
    content_index[element["id"]] = {"slide_index": None, "text_written": alt, "placed": True}


# --------------------------------------------------------------------------- build
def build(manifest, plan, out_path, assets_dir):
    by_id = {e["id"]: e for e in manifest["elements"]}
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = prs.slide_layouts[6]

    font_log, gradient_log = [], []
    content_index = {}
    images_cache_dir = assets_dir / "images"

    for splan in plan["slides"]:
        slide = prs.slides.add_slide(blank_layout)
        bg_element = by_id.get(splan.get("background_element_id")) if splan.get("background_element_id") else None
        apply_background(slide, prs, bg_element, gradient_log)

        title_font = None
        if splan.get("title_element_id"):
            title_font = (by_id[splan["title_element_id"]].get("style") or {}).get("font_family")
        if splan.get("title"):
            add_title(slide, splan["title"], title_font, font_log)
            if splan.get("title_element_id"):
                content_index[splan["title_element_id"]] = {"slide_index": None, "text_written": splan["title"]}

        # image_feature gets a genuine two-column split (image left, text right) so the
        # two never occupy the same box -- every other layout uses the full body width.
        is_split = splan["layout"] == "image_feature" and any(
            by_id[eid]["type"] == "image" for eid in splan["element_ids"]
        ) and any(by_id[eid]["type"] in CONTENT_TYPES for eid in splan["element_ids"])
        col_gap = 0.4
        col_w = (SLIDE_WIDTH_IN - 2 * MARGIN_IN - col_gap) / 2
        image_region = (MARGIN_IN, BODY_TOP_IN, col_w, BODY_HEIGHT_IN) if is_split else None
        text_region = (MARGIN_IN + col_w + col_gap, BODY_TOP_IN, col_w, BODY_HEIGHT_IN) if is_split else None

        text_ids, text_lines, text_font = [], [], None
        for eid in splan["element_ids"]:
            el = by_id[eid]
            if el["type"] == "table":
                add_table(slide, el.get("rows", []), eid, content_index, gradient_log)
            elif el["type"] == "image":
                add_image(slide, el, images_cache_dir, content_index, gradient_log, region=image_region)
            elif el["type"] in CONTENT_TYPES:
                text_ids.append(eid)
                text_lines.append(el["text"])
                text_font = text_font or (el.get("style") or {}).get("font_family")

        if text_lines:
            add_body_textblock(slide, text_lines, content_index, text_ids, text_font, font_log,
                                bullets=(splan["layout"] == "bullet_stack"), region=text_region)

        for eid in list(content_index):
            if content_index[eid]["slide_index"] is None:
                content_index[eid]["slide_index"] = splan["index"]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {
        "output": str(out_path),
        "slide_count": len(plan["slides"]),
        "font_substitutions": font_log,
        "gradient_notes": gradient_log,
        "content_index": content_index,
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("manifest", type=Path)
    ap.add_argument("--plan", type=Path, default=None)
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    out_path = args.out or args.manifest.with_suffix("").with_suffix(".pptx")
    assets_dir = args.manifest.parent / (args.manifest.stem.replace(".manifest", "") + "_build_assets")

    if args.plan and args.plan.exists():
        plan = json.loads(args.plan.read_text(encoding="utf-8"))
    else:
        plan = default_plan(manifest)
        plan_path = args.manifest.with_suffix("").with_suffix(".slide_plan.json")
        plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
        print("no --plan given; wrote default heuristic plan to " + str(plan_path), file=sys.stderr)

    report = build(manifest, plan, out_path, assets_dir)
    report_path = out_path.with_suffix(".build_report.json")
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print("pptx: " + str(out_path) + "  (" + str(report["slide_count"]) + " slides)")
    print("build report: " + str(report_path))


if __name__ == "__main__":
    main()
